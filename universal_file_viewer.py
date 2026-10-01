# =====================================================================================
#  UNIVERSAL FILE VIEWER  -  Google Colab notebook cells  (beginner-friendly edition)
# -------------------------------------------------------------------------------------
#  WHAT THIS PROJECT DOES
#    You upload ANY file (pdf, image, word, excel, csv, code, zip, mp3, unknown...).
#    The program reads it as text AND shows it as images, charts, diagrams, tables or
#    media players wherever possible. It works in two places:
#       (a) directly inside the Colab output   -> CELL 5
#       (b) as a Streamlit web app              -> CELL 6
#
#  HOW TO USE THIS FILE
#    Create one Colab cell per "CELL N" block and paste that block into it.
#    For CELL 2, 3, 4a and 4b the VERY FIRST line of the cell must be the
#    %%writefile line (Colab "magics" only work on the first line of a cell).
#
#  FILES THIS NOTEBOOK CREATES (push these 4 to GitHub to deploy on Streamlit Cloud):
#     file_analyzer.py   -> the "brain": understands every file type (shared code)
#     app.py             -> the Streamlit web page (the "face")
#     requirements.txt   -> list of Python packages to install
#     packages.txt       -> system program needed for OCR (tesseract)
#
#  ─────────────── MINI PYTHON GLOSSARY (read once, refer back anytime) ───────────────
#   variable        a name that stores a value:            age = 25
#   str             text:                                   "hello"
#   bytes           raw file data (0-255 numbers):          b"\x89PNG..."
#   list            ordered collection:                     [1, 2, 3]
#   dict            key -> value pairs:                     {"name": "Saavan", "age": 25}
#   tuple           fixed collection:                       (1, "a")
#   set             collection without duplicates:          {"png", "jpg"}
#   function        reusable block of code:                 def add(a, b): return a + b
#   import          load code written by others (a "library" / "module")
#   f-string        text with values inside:                f"Hi {name}"  -> "Hi Saavan"
#   if/elif/else    choose what to do based on a condition
#   for loop        repeat for each item:                   for x in [1, 2]: print(x)
#   try/except      "try this; if it crashes, do this instead" (error handling)
#   with ...:       open something and close it automatically afterwards
#   list comprehension  a short way to build a list:        [x * 2 for x in nums]
#   None            "nothing / no value"
#   True / False    booleans (yes / no)
# =====================================================================================


# ╔══════════════════════════════ CELL 1 : install ══════════════════════════════╗
# Lines starting with "!" are NOT Python - they run as terminal (shell) commands in Colab.
# apt-get installs system programs:  tesseract-ocr = reads text from images (OCR),
#                                    graphviz      = draws diagrams in Colab.
# "> /dev/null" just hides the long installation messages.
!apt-get -qq install -y tesseract-ocr graphviz > /dev/null

# pip installs Python libraries. "-q" = quiet mode (less output).
#   streamlit     -> builds the web app            pymupdf     -> reads PDF files
#   python-docx   -> reads Word .docx files        python-pptx -> reads PowerPoint .pptx
#   openpyxl      -> lets pandas read Excel        chardet     -> guesses text encoding
#   pytesseract   -> Python bridge to tesseract    graphviz    -> Python bridge to graphviz
!pip install -q streamlit pymupdf python-docx python-pptx openpyxl chardet pytesseract graphviz

# ╔══ NEXT: CELL 2 → create a NEW cell, its first line must be: %%writefile file_analyzer.py ══╗



%%writefile file_analyzer.py
# ↑ "%%writefile" is a Colab magic: instead of RUNNING this cell, it SAVES everything
#   below into a file called file_analyzer.py. app.py and CELL 5 will "import" it.

"""
file_analyzer.py  -  the "brain" of the project
-----------------------------------------------
(Text between triple quotes at the top of a file is called a "docstring":
 a description for humans. Python ignores it when running.)

BIG IDEA:
    analyze() never DRAWS anything on screen. It only RETURNS a list of
    "Output" objects (little boxes saying "show this text", "show this image"...).
    Each front-end then decides HOW to show them:
        - app.py        -> shows them in a Streamlit web page
        - Colab CELL 5  -> shows them in the notebook output
    So the same logic powers both, and you only write it once.
"""

# ---------------------------------------------------------------------
# IMPORTS: loading tools other people wrote, so we don't reinvent them.
# ---------------------------------------------------------------------
import ast              # reads Python code as a "tree" so we can find functions/classes
import base64           # turns bytes into safe text (used to embed images in HTML)
import collections      # Counter (counts things) and OrderedDict (dict that keeps order)
import html             # html.escape() makes text safe to put inside HTML
import io               # io.BytesIO lets us treat bytes in memory like a file
import json             # reads/writes JSON data
import mimetypes        # guesses a file type from its name (e.g. "a.png" -> "image/png")
import os               # os.path helpers (file names, extensions)
import re               # "regular expressions": powerful text search patterns
import zipfile          # opens .zip files (and docx/pptx/xlsx, which are zips inside!)
import xml.etree.ElementTree as ET   # reads XML. "as ET" gives it a short nickname
from dataclasses import dataclass, field   # quick way to create simple data classes

import numpy as np              # fast maths on lists of numbers ("np" = common nickname)
import pandas as pd             # tables (DataFrames), like Excel inside Python
import matplotlib.pyplot as plt # draws charts
from PIL import Image, ExifTags # Pillow: opens and edits images

# ---------------------------------------------------------------------
# OPTIONAL imports. "try / except ImportError" means:
#   try to import the library; if it isn't installed, set it to None
#   instead of crashing. Later we check "if fitz is None" before using it.
# ---------------------------------------------------------------------
try:
    import fitz                        # PyMuPDF -> PDF text, page images, embedded images
except ImportError:
    fitz = None
try:
    import docx                        # python-docx -> Word (.docx)
except ImportError:
    docx = None
try:
    from pptx import Presentation      # python-pptx -> PowerPoint (.pptx)
except ImportError:
    Presentation = None
try:
    import chardet                     # guesses which encoding a text file uses
except ImportError:
    chardet = None
try:
    import pytesseract                 # OCR: reads text inside images / scanned PDFs
except ImportError:
    pytesseract = None


# =====================================================================
# SETTINGS
# Names in CAPITAL_LETTERS are "constants": values we set once and don't
# change while running. Keeping them at the top makes them easy to tweak.
# =====================================================================
MAX_TEXT_CHARS = 50_000        # "_" in numbers is just for readability (= 50000)
MAX_TABLE_ROWS = 1_000         # show at most 1000 rows of a table
MAX_PDF_PAGES = 5              # PDF pages turned into pictures (text still covers ALL pages)
MAX_EMBEDDED_IMAGES = 12       # max pictures pulled out of PDFs / Word / zips
MAX_TREE_NODES = 150           # max boxes in a generated diagram (keeps it readable)
MAX_MEDIA_MB = 40              # audio/video bigger than this isn't embedded in Colab

# A SET of image extensions. Sets are great for "is X one of these?" checks.
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tif", ".tiff", ".ico"}

# DICTS: extension -> "MIME type" (the official name browsers use for a file type)
AUDIO_EXT = {".mp3": "audio/mpeg", ".wav": "audio/wav", ".ogg": "audio/ogg",
             ".m4a": "audio/mp4", ".flac": "audio/flac", ".aac": "audio/aac"}
VIDEO_EXT = {".mp4": "video/mp4", ".webm": "video/webm", ".mov": "video/quicktime",
             ".m4v": "video/mp4", ".ogv": "video/ogg"}

GRAPHVIZ_EXT = {".dot", ".gv"}        # Graphviz diagram files
MERMAID_EXT = {".mmd", ".mermaid"}    # Mermaid diagram files

# extension -> programming language name (used for colourful code display)
CODE_LANG = {
    ".py": "python", ".js": "javascript", ".ts": "typescript", ".jsx": "javascript",
    ".java": "java", ".c": "c", ".h": "c", ".cpp": "cpp", ".cs": "csharp", ".go": "go",
    ".rs": "rust", ".rb": "ruby", ".php": "php", ".sh": "bash", ".sql": "sql",
    ".css": "css", ".yaml": "yaml", ".yml": "yaml", ".r": "r", ".kt": "kotlin",
    ".swift": "swift", ".toml": "toml", ".ini": "ini", ".json": "json", ".xml": "xml",
    ".html": "html", ".htm": "html", ".md": "markdown",
}

# Common little words we DON'T want in the "most frequent words" chart.
# """...""" is a multi-line string; .split() cuts it into a list of words; set() removes duplicates.
STOPWORDS = set("""the and for are but not you all any can had her was one our out has have
him his how its may new now old see two way who did get let put say she too use with this that
from they will would there their what about which when your into than them then these some were
been also just more only over such very like here where while each other most must should could""".split())

# Which tab/section each kind of output goes into (used to build the tabs)
SECTION_OF = {
    "info": "Overview", "note": "Overview", "warning": "Overview",
    "text": "Text", "code": "Text", "markdown": "Text",
    "image": "Visuals", "svg": "Visuals",
    "chart": "Charts",
    "graphviz": "Diagrams", "mermaid": "Diagrams",
    "table": "Data", "json": "Data",
    "html": "Preview", "audio": "Preview", "video": "Preview",
}
SECTION_ORDER = ["Overview", "Text", "Visuals", "Charts", "Diagrams", "Data", "Preview"]

# Graphviz "DOT" is a small language for describing diagrams, e.g.  a -> b;
# This header sets the look of every diagram we generate (left-to-right, rounded boxes).
DOT_HEADER = ('digraph G {\n  rankdir=LR;\n'
              '  node [shape=box, style="rounded,filled", fillcolor="#EEF3FB", fontname="Helvetica", fontsize=10];\n'
              '  edge [color="#888888"];')


# ---------------------------------------------------------------------
# @dataclass is a "decorator": it automatically writes the boring setup
# code for a class that just holds data. Output("text", "hello", "Title")
# creates an object with .kind, .data, .title and .extra.
# ---------------------------------------------------------------------
@dataclass
class Output:
    """One thing to display.
    kind  : what it is -> info | note | warning | text | code | markdown | image | svg |
            chart | graphviz | mermaid | table | json | html | audio | video
    data  : the content (text, bytes, a table, a chart, a dict ...)
    title : heading shown above it
    extra : extra settings, e.g. {"language": "python"} or {"mime": "audio/mpeg"}
    """
    kind: str                     # ": str" is a "type hint" - a note saying what type is expected
    data: object = None           # "= None" is a default value (used if you don't pass one)
    title: str = ""
    extra: dict = field(default_factory=dict)   # gives every Output its OWN empty dict


def group_outputs(outputs):
    """Sort outputs into sections (tabs), keeping their order inside each section."""
    # OrderedDict remembers insertion order -> tabs appear in SECTION_ORDER order.
    # (s, []) for s in SECTION_ORDER  creates  ("Overview", []), ("Text", []), ...
    groups = collections.OrderedDict((s, []) for s in SECTION_ORDER)
    for o in outputs:
        # .get(key, default) returns the default if the key isn't in the dict
        section = SECTION_OF.get(o.kind, "Overview")
        groups[section].append(o)        # .append adds an item to the end of a list
    return groups


def capabilities():
    """Report which optional features are available (shown in the app's sidebar)."""
    ocr_ok = False
    if pytesseract is not None:
        try:
            pytesseract.get_tesseract_version()   # fails if the tesseract PROGRAM is missing
            ocr_ok = True
        except Exception:
            pass                                   # "pass" = do nothing
    # "is not None" -> True if the library was imported successfully
    return {"PDF (PyMuPDF)": fitz is not None, "Word (python-docx)": docx is not None,
            "PowerPoint (python-pptx)": Presentation is not None,
            "Encoding detection (chardet)": chardet is not None, "OCR (tesseract)": ocr_ok}


# =====================================================================
# SMALL HELPER FUNCTIONS
# Small functions that each do ONE job are easier to read, test and reuse.
# =====================================================================
def get_ext(name):
    """'Report.PDF' -> '.pdf'   (splitext splits name and extension; lower() = small letters)"""
    return os.path.splitext(name)[1].lower()


def human_size(n):
    """1536 -> '1.5 KB'. Divides by 1024 until the number is small enough."""
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024:
            # {n:.1f} formats a number with 1 decimal place; {n:.0f} with none
            return f"{n:.0f} {unit}" if unit == "B" else f"{n:.1f} {unit}"
        n /= 1024                       # same as: n = n / 1024
    return f"{n:.1f} TB"


def q(text, maxlen=40):
    """Make text safe to use as a label inside a Graphviz diagram.
    Quotes and backslashes have special meaning in DOT, so we "escape" them with \\ ."""
    s = str(text).replace("\n", " ").replace("\\", "\\\\").replace('"', '\\"')
    if len(s) > maxlen:
        s = s[:maxlen - 1] + "…"        # s[:10] = first 10 characters ("slicing")
    return f'"{s}"'


def truncate(text, limit=MAX_TEXT_CHARS):
    """Cut very long text so the page stays fast, and say that we cut it."""
    if len(text) <= limit:
        return text
    return text[:limit] + f"\n\n... [truncated: showing {limit:,} of {len(text):,} characters]"
    #                                                   {x:,} adds commas: 50000 -> 50,000


def unique_columns(cols):
    """Tables can't have two columns with the same name -> add _2, _3 ... to repeats."""
    seen = collections.Counter()        # Counter = a dict that counts: seen["Name"] += 1
    out = []
    for c in cols:
        c = str(c).strip() or "column"  # "x or y" -> x if x isn't empty, otherwise y
        seen[c] += 1
        out.append(c if seen[c] == 1 else f"{c}_{seen[c]}")
    return out


def sniff(data):
    """Guess the REAL file type from its first bytes ("magic numbers"), ignoring the name.
    Every PNG starts with b'\\x89PNG', every PDF with b'%PDF', every ZIP with b'PK'..."""
    # data[:4] = first 4 bytes;  .startswith() checks the beginning
    if data.startswith(b"\x89PNG") or data.startswith(b"\xff\xd8\xff") or data[:4] == b"GIF8":
        return "image"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image"
    if data[:4] == b"RIFF" and data[8:12] == b"WAVE":
        return "audio"
    if data.startswith(b"ID3") or data[:4] == b"fLaC" or data[:4] == b"OggS":
        return "audio"
    if data[4:8] == b"ftyp":
        return "video"
    if data.startswith(b"%PDF"):
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        return "zip"
    return None                         # unknown


def looks_binary(data):
    """Is this file binary (like an .exe) or normal text?
    Text files almost never contain the 'null' byte (0) or many control characters."""
    sample = data[:8192]                # only check the first 8 KB (fast)
    if not sample:                      # empty file
        return False
    if b"\x00" in sample:
        return True
    # sum(1 for ...) counts how many bytes are "control characters"
    control = sum(1 for b in sample if b < 9 or 13 < b < 32)
    return control / len(sample) > 0.10     # more than 10% control chars -> binary


def decode_text(data):
    """bytes -> str. Files store text using an "encoding" (rules for turning
    letters into bytes). We try UTF-8 first (most common), then ask chardet
    to guess, and finally latin-1 (which never fails).
    Returns TWO values (a tuple): the text and the encoding used."""
    try:
        return data.decode("utf-8-sig"), "UTF-8"
    except UnicodeDecodeError:
        pass
    if chardet is not None:
        guess = chardet.detect(data[:200_000]).get("encoding")
        if guess:
            try:
                return data.decode(guess), guess
            except (UnicodeDecodeError, LookupError):
                pass
    return data.decode("latin-1"), "latin-1 (fallback)"


def image_bytes_for_display(img, raw=None):
    """Browsers can show PNG/JPEG/GIF/WEBP directly. Other formats (BMP, TIFF, ICO...)
    are converted to PNG first. GIFs are kept as-is so animations still play."""
    if raw is not None and img.format in ("PNG", "JPEG", "GIF", "WEBP"):
        return raw
    buf = io.BytesIO()                  # an empty "file" that lives in memory
    try:
        frame = img if img.mode in ("RGB", "RGBA", "L") else img.convert("RGBA")
    except Exception:
        frame = img.convert("RGB")
    frame.save(buf, format="PNG")       # save the picture INTO the memory file
    return buf.getvalue()               # get the bytes back out


def raw_image_to_display(raw):
    """Any image bytes -> displayable bytes, or None if Pillow can't read it
    (some Office files contain EMF/WMF pictures that Pillow doesn't support)."""
    try:
        img = Image.open(io.BytesIO(raw))
        img.load()                      # actually read the pixels (finds broken files early)
        return image_bytes_for_display(img, raw)
    except Exception:
        return None


def ocr_text(img):
    """OCR = Optical Character Recognition: reading text that is inside a picture."""
    if pytesseract is None:
        return None
    try:
        text = pytesseract.image_to_string(img.convert("RGB")).strip()   # strip() removes spaces/newlines at the ends
        return text or None             # empty text -> None
    except Exception:
        return None


def mermaid_html(src, height=450):
    """Mermaid is a diagram language that draws IN THE BROWSER using JavaScript.
    This returns a little HTML page that loads Mermaid and draws `src`.
    Note: {{ and }} inside an f-string mean literal { and } characters."""
    return f"""
<div class="mermaid" style="min-height:{height - 50}px">{html.escape(src)}</div>
<script type="module">
  import mermaid from 'https://cdn.jsdelivr.net/npm/mermaid@11/dist/mermaid.esm.min.mjs';
  mermaid.initialize({{ startOnLoad: true }});
</script>"""


def svg_html(svg_text):
    """Show an SVG image by embedding it as base64 inside an <img> tag."""
    b64 = base64.b64encode(svg_text.encode("utf-8")).decode()
    return f'<img src="data:image/svg+xml;base64,{b64}" style="max-width:100%;max-height:600px">'


def bar_chart(labels, values, title, xlabel="", horizontal=True):
    """Draw a bar chart with matplotlib and RETURN the figure (we don't show it here)."""
    labels = [str(l)[:30] for l in labels]          # list comprehension: shorten every label
    if horizontal:
        # figsize = (width, height) in inches; taller when there are more bars
        fig, ax = plt.subplots(figsize=(7, max(2.5, 0.35 * len(labels))))
        ax.barh(labels[::-1], values[::-1], color="#4C78A8")   # [::-1] reverses a list -> biggest on top
        ax.set_xlabel(xlabel)
    else:
        fig, ax = plt.subplots(figsize=(8, 3.5))
        ax.bar(labels, values, color="#4C78A8")
        ax.set_ylabel(xlabel)
        ax.tick_params(axis="x", rotation=45)        # tilt labels so they don't overlap
    ax.set_title(title)
    fig.tight_layout()                               # auto-fix spacing so nothing is cut off
    return fig


# =====================================================================
# TEXT  (every file ends up here in some form -> "consider it as text")
# =====================================================================
def text_outputs(text, title="Text content", language=None, word_chart=True):
    """For any text: statistics + the text itself + a word-frequency chart."""
    lines = text.splitlines()                        # split text into a list of lines
    outs = [Output("info", {
        "Characters": f"{len(text):,}",
        # re.findall(pattern, text) returns every match; \S+ = a run of non-space characters (a word)
        "Words": f"{len(re.findall(r'\S+', text)):,}",
        "Lines": f"{len(lines):,}",
        "Non-empty lines": f"{sum(1 for l in lines if l.strip()):,}",
    }, f"{title} – statistics")]

    # Code gets colourful "code" display; normal text gets plain "text" display
    outs.append(Output("code" if language else "text", truncate(text), title, {"language": language}))

    if word_chart:
        # [^\W\d_]{3,}  = words of 3+ LETTERS in any language (Hindi, Gujarati, English...)
        words = [w.lower() for w in re.findall(r"[^\W\d_]{3,}", text)]
        # Counter counts each word; .most_common(15) = top 15 as [(word, count), ...]
        counts = collections.Counter(w for w in words if w not in STOPWORDS).most_common(15)
        if len(counts) >= 3:
            labels, values = zip(*counts)            # [("a",5),("b",3)] -> ("a","b") and (5,3)
            outs.append(Output("chart", bar_chart(list(labels), list(values), "Most frequent words", "count"),
                               f"{title} – word frequency"))
    return outs


# Words that a Mermaid diagram usually starts with
MERMAID_START = ("graph ", "graph\n", "flowchart", "sequenceDiagram", "classDiagram", "stateDiagram",
                 "erDiagram", "gantt", "pie", "journey", "mindmap", "timeline", "gitGraph")


def diagram_outputs_from_text(text):
    """Find diagrams written as TEXT and turn them into real pictures:
       - ```mermaid ... ``` or ```dot ... ``` blocks (common in markdown files)
       - a whole file that IS a Graphviz / Mermaid diagram."""
    outs = []
    # (.*?) = capture anything, as little as possible. re.S lets "." match newlines too,
    # re.I = ignore upper/lower case.
    for lang, body in re.findall(r"```(mermaid|dot|graphviz)[^\n]*\n(.*?)```", text, re.S | re.I):
        kind = "mermaid" if lang.lower() == "mermaid" else "graphviz"
        outs.append(Output(kind, body.strip(), f"{lang} diagram found in text"))
    if not outs:
        s = text.lstrip()
        if re.match(r"(strict\s+)?(di)?graph\b", s, re.I) and "{" in s:
            outs.append(Output("graphviz", s, "Graphviz diagram"))
        elif s.startswith(MERMAID_START):           # startswith accepts a tuple of options
            outs.append(Output("mermaid", s, "Mermaid diagram"))
    return outs


def python_structure_dot(code):
    """Draw a Python file's structure as a diagram:
    classes, their methods, functions, and which function calls which."""
    try:
        tree = ast.parse(code)          # turn code into a tree of "nodes" Python understands
    except SyntaxError:                 # the code has errors -> no diagram
        return None

    # all top-level functions (isinstance checks "is this object of this type?")
    funcs = [n for n in tree.body if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef))]
    # dict comprehension: {"main": "f0", "helper": "f1", ...} -> short ids for the diagram
    func_ids = {f.name: f"f{i}" for i, f in enumerate(funcs)}   # enumerate gives (index, item)

    lines = [DOT_HEADER, '  module [label="module", shape=folder, fillcolor="#D6EAF8"];']
    imports, count = [], 0              # two variables set in one line

    for node in tree.body:              # tree.body = the top-level statements of the file
        if isinstance(node, ast.Import):                       # import x
            imports += [a.name for a in node.names]            # += adds items to the list
        elif isinstance(node, ast.ImportFrom):                 # from x import y
            imports += [f"{node.module}.{a.name}" for a in node.names]
        elif isinstance(node, ast.ClassDef):                   # class X:
            cid = f"c{count}"; count += 1                      # ";" puts two statements on one line
            lines.append(f'  {cid} [label={q("class " + node.name)}, fillcolor="#FDEBD0"];')
            lines.append(f"  module -> {cid};")                # "->" draws an arrow in DOT
            for item in node.body:                             # methods inside the class
                if isinstance(item, (ast.FunctionDef, ast.AsyncFunctionDef)) and count < MAX_TREE_NODES:
                    mid = f"m{count}"; count += 1
                    lines.append(f"  {mid} [label={q(item.name + '()')}];")
                    lines.append(f"  {cid} -> {mid};")
        elif isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):   # def x():
            count += 1
            lines.append(f"  {func_ids[node.name]} [label={q(node.name + '()')}];")
            lines.append(f"  module -> {func_ids[node.name]};")

    # Find "function A calls function B" -> dashed red arrows.
    calls = set()                       # a set ignores duplicates automatically
    for f in funcs:
        for sub in ast.walk(f):         # ast.walk visits EVERY node inside the function
            if isinstance(sub, ast.Call) and isinstance(sub.func, ast.Name) \
                    and sub.func.id in func_ids and sub.func.id != f.name:
                calls.add((f.name, sub.func.id))
    for a, b in calls:
        lines.append(f'  {func_ids[a]} -> {func_ids[b]} [style=dashed, color="#C0392B", label="calls", fontsize=8];')

    if imports:
        label = "imports: " + ", ".join(imports[:12]) + (" …" if len(imports) > 12 else "")
        lines.append(f'  imports [label={q(label, 120)}, shape=note, fillcolor="#E8F8F5"];')
        lines.append("  module -> imports;")
    if count == 0 and not imports:      # nothing interesting found
        return None
    lines.append("}")
    return "\n".join(lines)             # join the list of lines into one string


def headings_dot(title, headings):
    """headings = [(level, text), ...]  e.g. [(1, "Intro"), (2, "Goals")]
    -> a tree diagram where level-2 headings hang under the level-1 above them."""
    lines = [DOT_HEADER, f'  root [label={q(title)}, shape=folder, fillcolor="#D6EAF8"];']
    stack = [(0, "root")]               # a "stack" remembers the current chain of parents
    for i, (level, text) in enumerate(headings[:MAX_TREE_NODES]):
        nid = f"h{i}"
        lines.append(f"  {nid} [label={q(text)}];")
        # go back up until we find a parent with a SMALLER level
        while len(stack) > 1 and stack[-1][0] >= level:    # stack[-1] = last item
            stack.pop()                                     # pop() removes the last item
        lines.append(f"  {stack[-1][1]} -> {nid};")
        stack.append((level, nid))
    lines.append("}")
    return "\n".join(lines)


def analyze_text(data, ext, name):
    """Plain text and code files."""
    text, encoding = decode_text(data)          # "unpacking": a function returning 2 values
    language = CODE_LANG.get(ext)               # None if it isn't a known code file
    outs = [Output("note", f"Read as text using the {encoding} encoding.")]
    outs += text_outputs(text, "File content", language, word_chart=language in (None, "markdown"))
    if ext == ".py":
        dot = python_structure_dot(text)
        if dot:                                 # None counts as False in an if
            outs.append(Output("graphviz", dot, "Code structure (classes, functions, calls)"))
    outs += diagram_outputs_from_text(text)
    return outs


# =====================================================================
# IMAGES
# =====================================================================
def analyze_image(data):
    img = Image.open(io.BytesIO(data))          # open the image from memory bytes
    img.load()
    info = {"Format": img.format, "Dimensions": f"{img.width} × {img.height} px",
            "Colour mode": img.mode,
            # getattr(obj, "name", default) -> obj.name if it exists, otherwise default
            "Frames": getattr(img, "n_frames", 1)}
    try:
        # EXIF = hidden photo details (camera model, date taken...). TAGS turns numbers into names.
        for tag_id, value in list(img.getexif().items())[:15]:
            info[f"EXIF {ExifTags.TAGS.get(tag_id, tag_id)}"] = str(value)[:60]
    except Exception:
        pass
    outs = [Output("info", info, "Image details"),
            Output("image", image_bytes_for_display(img, data), "Image preview")]

    # Colour histogram: for each colour, how many pixels have each brightness (0 dark - 255 bright)
    try:
        hist = img.convert("RGB").histogram()   # 768 numbers: 256 red, then 256 green, then 256 blue
        fig, ax = plt.subplots(figsize=(7, 3))
        for i, colour in enumerate(("red", "green", "blue")):
            ax.plot(range(256), hist[i * 256:(i + 1) * 256], color=colour, label=colour)
        ax.set_title("Colour histogram"); ax.set_xlabel("pixel value (0–255)"); ax.legend()
        fig.tight_layout()
        outs.append(Output("chart", fig, "Colour distribution"))
    except Exception:
        pass

    # "Consider it as text": try to read any words inside the picture
    text = ocr_text(img)
    if text:
        outs += text_outputs(text, "Text found in image (OCR)")
    else:
        outs.append(Output("note", "No readable text detected in this image." if pytesseract
                           else "OCR not available - install tesseract + pytesseract to read text in images."))
    return outs


def analyze_svg(data):
    """SVG images are actually TEXT (XML) describing shapes -> show the picture AND the source."""
    text, _ = decode_text(data)                 # "_" = a value we don't need
    return [Output("svg", text, "SVG preview")] + text_outputs(text, "SVG source", "xml", word_chart=False)


# =====================================================================
# PDF
# =====================================================================
def analyze_pdf(data):
    if fitz is None:
        return [Output("warning", "PyMuPDF is not installed, so the PDF is shown as raw data.")] + analyze_binary(data)

    doc = fitz.open(stream=data, filetype="pdf")
    meta = doc.metadata or {}                   # "or {}" -> use an empty dict if metadata is None
    outs = [Output("info", {"Pages": doc.page_count, "Title": meta.get("title") or "–",
                            "Author": meta.get("author") or "–", "Created with": meta.get("producer") or "–"},
                   "PDF details")]

    page_texts, rendered = [], []
    for i, page in enumerate(doc):              # loop over pages; i = 0, 1, 2 ...
        page_texts.append(page.get_text())
        if i < MAX_PDF_PAGES:
            # Matrix(1.5, 1.5) = zoom 150% so the page picture is sharper
            png = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).tobytes("png")
            rendered.append(png)
            outs.append(Output("image", png, f"Page {i + 1}"))
    if doc.page_count > MAX_PDF_PAGES:
        outs.append(Output("note", f"First {MAX_PDF_PAGES} pages shown as images; the text covers all {doc.page_count} pages."))

    # A scanned PDF is just photos of pages (no real text) -> read it with OCR instead
    if sum(len(t.strip()) for t in page_texts) < 50 and pytesseract is not None:
        ocr = [ocr_text(Image.open(io.BytesIO(p))) or "" for p in rendered]
        if any(ocr):                            # any() -> True if at least one item has text
            page_texts = ocr
            outs.append(Output("note", "This looks like a scanned PDF - text was read with OCR."))

    full_text = "\n\n".join(f"--- Page {i + 1} ---\n{t}" for i, t in enumerate(page_texts))
    outs += text_outputs(full_text, "Extracted text")

    if 1 < len(page_texts) <= 60:               # chained comparison: between 2 and 60 pages
        counts = [len(t.split()) for t in page_texts]
        outs.append(Output("chart", bar_chart([str(i + 1) for i in range(len(counts))], counts,
                                              "Words per page", "words", horizontal=False), "Words per page"))

    # Pull out pictures stored inside the PDF. "xref" is the PDF's internal id for an image;
    # we remember seen ids so the same logo on every page is only shown once.
    seen, found = set(), 0
    for page in doc:
        for im in page.get_images(full=True):
            xref = im[0]
            if xref in seen or found >= MAX_EMBEDDED_IMAGES:
                continue                        # "continue" = skip to the next loop item
            seen.add(xref)
            try:
                shown = raw_image_to_display(doc.extract_image(xref)["image"])
            except Exception:
                shown = None
            if shown:
                found += 1
                outs.append(Output("image", shown, f"Embedded image {found} (page {page.number + 1})"))
    return outs


# =====================================================================
# OFFICE FILES (docx / pptx / xlsx)
# Fun fact: these are secretly ZIP files full of XML + a "media" folder of pictures!
# =====================================================================
def zip_images(data, prefix):
    """Pull the pictures out of an Office file's media folder (e.g. 'word/media/')."""
    outs = []
    # "with" opens the zip and closes it automatically when the block ends
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = [n for n in z.namelist() if n.startswith(prefix) and get_ext(n) in IMAGE_EXT]
        for n in names[:MAX_EMBEDDED_IMAGES]:
            shown = raw_image_to_display(z.read(n))
            if shown:
                outs.append(Output("image", shown, f"Embedded image: {os.path.basename(n)}"))
    return outs


def analyze_docx(data):
    if docx is None:
        return [Output("warning", "python-docx is not installed.")] + analyze_zip(data)
    d = docx.Document(io.BytesIO(data))
    md, headings, plain = [], [], []            # md = markdown lines (keeps headings/bullets)
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:                               # skip empty paragraphs
            continue
        plain.append(t)
        style = (p.style.name if p.style is not None else "").lower()   # e.g. "heading 2"
        if style.startswith("heading") or style == "title":
            digits = re.findall(r"\d+", style)  # find the number in "heading 2"
            level = int(digits[0]) if digits else 1
            md.append("#" * min(level + 1, 6) + " " + t)   # "#" * 3 = "###" (markdown heading)
            headings.append((level, t))
        elif "list" in style:
            md.append(f"- {t}")                 # markdown bullet
        else:
            md.append(t)

    outs = [Output("info", {"Paragraphs": len(plain), "Headings": len(headings), "Tables": len(d.tables)},
                   "Word document details"),
            Output("markdown", "\n\n".join(md) or "_(no text)_", "Document (formatted)")]
    outs += text_outputs("\n".join(plain), "Document text")
    if headings:
        outs.append(Output("graphviz", headings_dot("Document", headings), "Document outline"))
    for i, t in enumerate(d.tables, 1):         # enumerate(..., 1) starts counting at 1
        rows = [[c.text for c in r.cells] for r in t.rows]     # nested list comprehension
        if rows:
            # first row = column names, the rest = data
            df = pd.DataFrame(rows[1:], columns=unique_columns(rows[0])) if len(rows) > 1 else pd.DataFrame(rows)
            outs.append(Output("table", df, f"Table {i}"))
    outs += zip_images(data, "word/media/")
    return outs


def analyze_pptx(data):
    if Presentation is None:
        return [Output("warning", "python-pptx is not installed.")] + analyze_zip(data)
    prs = Presentation(io.BytesIO(data))
    md, all_text, titles, words_per_slide = [], [], [], []
    for i, slide in enumerate(prs.slides, 1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None and title_shape.text.strip() else f"Slide {i}"
        texts = []
        for shape in slide.shapes:              # every box/picture on the slide
            if shape.has_text_frame:
                texts += [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]
        notes = ""
        if slide.has_notes_slide:               # speaker notes under the slide
            frame = slide.notes_slide.notes_text_frame
            notes = frame.text.strip() if frame is not None else ""
        body = "\n".join(f"- {t}" for t in texts if t != title)
        md.append(f"### Slide {i}: {title}\n{body}" + (f"\n\n> Notes: {notes}" if notes else ""))
        all_text += [title] + texts + ([notes] if notes else [])
        titles.append((1, f"{i}. {title}"))
        words_per_slide.append(len(" ".join(texts).split()))

    outs = [Output("info", {"Slides": len(prs.slides)}, "Presentation details"),
            Output("markdown", "\n\n".join(md), "Slides (text)")]
    outs += text_outputs("\n".join(all_text), "Presentation text")
    if titles:
        outs.append(Output("graphviz", headings_dot("Presentation", titles), "Slide outline"))
    if len(words_per_slide) > 1:
        outs.append(Output("chart", bar_chart([str(i + 1) for i in range(len(words_per_slide))], words_per_slide,
                                              "Words per slide", "words", horizontal=False), "Words per slide"))
    outs += zip_images(data, "ppt/media/")
    return outs


# =====================================================================
# TABLES (csv / tsv / excel / json lists)
# =====================================================================
def safe_nunique(series):
    """Number of different values in a column (-1 if it contains lists, which can't be counted)."""
    try:
        return series.nunique()
    except TypeError:
        return -1


def dataframe_outputs(df, title):
    """For any table (DataFrame): summary, the data, statistics and automatic charts."""
    df = df.copy()                              # work on a copy so we don't change the original
    df.columns = unique_columns(df.columns)
    num = df.select_dtypes(include="number")    # only the numeric columns
    outs = [Output("info", {"Rows": f"{len(df):,}", "Columns": len(df.columns),
                            "Numeric columns": len(num.columns),
                            # isna() marks empty cells True; .sum().sum() counts them all
                            "Missing values": f"{int(df.isna().sum().sum()):,}"}, f"{title} – summary"),
            Output("table", df.head(MAX_TABLE_ROWS),          # head(n) = first n rows
                   title + (f" (first {MAX_TABLE_ROWS:,} rows)" if len(df) > MAX_TABLE_ROWS else ""))]

    if not num.empty:
        # describe() = count, mean, min, max... for every numeric column; .T flips rows/columns
        outs.append(Output("table", num.describe().T.round(3), f"{title} – statistics"))

        # 1) Histograms: how values are spread, for up to 4 numeric columns side by side
        cols = list(num.columns[:4])
        fig, axes = plt.subplots(1, len(cols), figsize=(3.6 * len(cols), 3))
        for ax, c in zip(np.atleast_1d(axes), cols):     # zip pairs items: (ax1, col1), (ax2, col2)...
            ax.hist(num[c].dropna(), bins=20, color="#4C78A8")   # dropna() skips empty cells
            ax.set_title(str(c)[:25], fontsize=9)
        fig.suptitle("Distributions"); fig.tight_layout()
        outs.append(Output("chart", fig, f"{title} – distributions"))

        # 2) Line chart: values row by row (shows trends if rows are in time order)
        if len(df) > 1:
            fig, ax = plt.subplots(figsize=(8, 3.2))
            for c in num.columns[:3]:
                ax.plot(num[c].values[:5000], label=str(c)[:25])
            ax.set_title("Values by row"); ax.set_xlabel("row"); ax.legend(); fig.tight_layout()
            outs.append(Output("chart", fig, f"{title} – trend"))

        # 3) Correlation heatmap: do columns go up/down together? (+1 = together, -1 = opposite)
        if len(num.columns) >= 2:
            corr = num.iloc[:, :12].corr()               # iloc[:, :12] = all rows, first 12 columns
            fig, ax = plt.subplots(figsize=(5.5, 4.5))
            im = ax.imshow(corr, cmap="RdBu_r", vmin=-1, vmax=1)
            ax.set_xticks(range(len(corr))); ax.set_xticklabels([str(c)[:12] for c in corr.columns], rotation=45, ha="right")
            ax.set_yticks(range(len(corr))); ax.set_yticklabels([str(c)[:12] for c in corr.columns])
            fig.colorbar(im); ax.set_title("Correlation between numeric columns"); fig.tight_layout()
            outs.append(Output("chart", fig, f"{title} – correlation"))

    # 4) Category chart: for the first text column with a few repeated values (e.g. "City")
    cat_cols = [c for c in df.columns if c not in num.columns and 1 < safe_nunique(df[c]) <= 30]
    if cat_cols:
        c = cat_cols[0]
        vc = df[c].astype(str).value_counts().head(15)   # value_counts = how often each value appears
        outs.append(Output("chart", bar_chart(list(vc.index), list(vc.values), f"Counts of '{c}'", "count"),
                           f"{title} – categories"))
    return outs


def analyze_csv(data, ext):
    text, _ = decode_text(data)
    # sep=None + engine="python" -> pandas auto-detects the separator ( , ; | tab )
    sep = "\t" if ext == ".tsv" else None
    df = pd.read_csv(io.StringIO(text), sep=sep, engine="python")   # StringIO = text "file" in memory
    return dataframe_outputs(df, "Table") + text_outputs(text, "Raw text", word_chart=False)


def analyze_excel(data):
    sheets = pd.read_excel(io.BytesIO(data), sheet_name=None)   # sheet_name=None -> ALL sheets as a dict
    outs = [Output("info", {"Sheets": ", ".join(map(str, sheets.keys()))}, "Workbook details")]
    for name, df in sheets.items():             # .items() gives (key, value) pairs
        outs += dataframe_outputs(df, f"Sheet '{name}'")
        outs.append(Output("text", truncate(df.to_csv(index=False)), f"Sheet '{name}' as text"))
    return outs


# =====================================================================
# STRUCTURED TEXT (json / notebook / markdown / html / xml)
# =====================================================================
def json_tree_dot(obj):
    """Draw JSON as a tree. Uses RECURSION: the inner function add() calls itself
    for every child, so it works no matter how deeply the data is nested."""
    lines, count = [DOT_HEADER], [0]            # count is a list so the inner function can change it

    def add(o, label):                          # a function defined inside a function
        if count[0] >= MAX_TREE_NODES:
            return None
        nid = f"n{count[0]}"; count[0] += 1
        if isinstance(o, dict):
            text, items = f"{label} {{{len(o)}}}", list(o.items())[:20]     # dict -> "name {3}"
        elif isinstance(o, list):
            text, items = f"{label} [{len(o)}]", [(f"[{i}]", v) for i, v in enumerate(o[:10])]
        else:
            text, items = f"{label}: {o}", []   # a simple value has no children
        lines.append(f"  {nid} [label={q(text)}];")
        for k, v in items:
            cid = add(v, k)                     # ← recursion: add the child (and ITS children)
            if cid:
                lines.append(f"  {nid} -> {cid};")
        return nid

    add(obj, "root")
    lines.append("}")
    return "\n".join(lines)


def analyze_json(data):
    text, _ = decode_text(data)
    obj = json.loads(text)                      # JSON text -> Python dicts/lists
    outs = [Output("json", obj, "JSON data"), Output("graphviz", json_tree_dot(obj), "JSON structure")]
    # A list of records like [{"name": .., "age": ..}, ...] can become a table
    if isinstance(obj, list) and obj and all(isinstance(x, dict) for x in obj[:50]):
        outs += dataframe_outputs(pd.json_normalize(obj), "JSON records")
    return outs + text_outputs(text, "JSON source", "json", word_chart=False)


def analyze_notebook(data):
    """.ipynb notebooks are JSON: a list of cells (markdown or code) with saved outputs."""
    nb = json.loads(decode_text(data)[0])       # [0] = first returned value (the text)
    cells = nb.get("cells", [])
    kinds = collections.Counter(c.get("cell_type", "?") for c in cells)
    # {**dict} "unpacks" one dict into another
    outs = [Output("info", {"Cells": len(cells), **{f"{k} cells": v for k, v in kinds.items()}}, "Notebook details")]
    for i, cell in enumerate(cells[:60], 1):
        src = cell.get("source", "")
        src = "".join(src) if isinstance(src, list) else str(src)
        if cell.get("cell_type") == "markdown":
            outs.append(Output("markdown", src, f"Cell {i} (markdown)"))
        else:
            outs.append(Output("code", src, f"Cell {i} (code)", {"language": "python"}))
            for out in cell.get("outputs", []):
                d = out.get("data", {})
                if "image/png" in d:           # charts saved in the notebook (base64 text)
                    png = d["image/png"]
                    png = "".join(png) if isinstance(png, list) else png
                    outs.append(Output("image", base64.b64decode(png), f"Cell {i} output image"))
                elif "text/plain" in d or "text" in out:
                    t = d.get("text/plain", out.get("text", ""))
                    outs.append(Output("text", "".join(t) if isinstance(t, list) else str(t), f"Cell {i} output"))
    return outs


def analyze_markdown(data):
    text, _ = decode_text(data)
    outs = [Output("markdown", text, "Rendered markdown")]
    outs += text_outputs(text, "Markdown source", "markdown")
    # ^(#{1,6})\s+(.+)$ = lines starting with 1-6 "#" then the heading text. re.M = check every line.
    headings = [(len(m.group(1)), m.group(2).strip()) for m in re.finditer(r"^(#{1,6})\s+(.+)$", text, re.M)]
    if headings:
        outs.append(Output("graphviz", headings_dot("Document", headings), "Heading outline"))
    return outs + diagram_outputs_from_text(text)


def analyze_html(data):
    text, _ = decode_text(data)
    visible = re.sub(r"(?is)<(script|style).*?</\1>", " ", text)   # remove <script>/<style> blocks
    visible = html.unescape(re.sub(r"<[^>]+>", " ", visible))       # remove all <tags>; &amp; -> &
    visible = re.sub(r"[ \t]+", " ", re.sub(r"\n\s*\n+", "\n\n", visible)).strip()   # tidy spaces
    outs = [Output("info", {"Links": len(re.findall(r"<a\s", text, re.I)),
                            "Images": len(re.findall(r"<img\s", text, re.I))}, "HTML details"),
            Output("html", text, "Rendered page")]
    outs += text_outputs(visible, "Visible text")
    return outs + text_outputs(text, "HTML source", "html", word_chart=False)


def xml_tree_dot(root):
    """Draw XML as a tree (same recursion idea as json_tree_dot)."""
    lines, count = [DOT_HEADER], [0]

    def add(el):
        if count[0] >= MAX_TREE_NODES:
            return None
        nid = f"x{count[0]}"; count[0] += 1
        tag = el.tag.split("}")[-1]            # "{namespace}book" -> "book"
        txt = (el.text or "").strip()
        lines.append(f"  {nid} [label={q(tag + (': ' + txt if txt else ''))}];")
        for child in list(el)[:15]:
            cid = add(child)
            if cid:
                lines.append(f"  {nid} -> {cid};")
        return nid

    add(root)
    lines.append("}")
    return "\n".join(lines)


def analyze_xml(data):
    text, _ = decode_text(data)
    outs = []
    try:
        outs.append(Output("graphviz", xml_tree_dot(ET.fromstring(data)), "XML structure"))
    except ET.ParseError as e:                  # "as e" stores the error so we can show it
        outs.append(Output("warning", f"XML could not be parsed: {e}"))
    return outs + text_outputs(text, "XML source", "xml", word_chart=False)


# =====================================================================
# MEDIA, ARCHIVES, UNKNOWN BINARY
# =====================================================================
def analyze_media(data, ext, kind):
    """kind is "audio" or "video" -> show a player."""
    if kind == "audio":
        mime = AUDIO_EXT.get(ext, "audio/mpeg")
    else:
        mime = VIDEO_EXT.get(ext, "video/mp4")
    # kind.title() -> "Audio" (first letter capital)
    return [Output(kind, data, f"{kind.title()} player", {"mime": mime}),
            Output("note", f"{kind.title()} has no text layer to show; play it in the Preview tab.")]


def folder_tree_dot(paths):
    """['a/b.txt', 'a/c.png'] -> a folder tree diagram."""
    nodes = {"": "root"}                        # path -> diagram id
    lines = [DOT_HEADER, '  root [label="(archive)", shape=folder, fillcolor="#D6EAF8"];']
    for p in paths:
        parts = [x for x in p.strip("/").split("/") if x]     # "a/b/c.txt" -> ["a", "b", "c.txt"]
        for depth in range(1, len(parts) + 1):
            key = "/".join(parts[:depth])       # "a", then "a/b", then "a/b/c.txt"
            if key in nodes or len(nodes) > MAX_TREE_NODES:
                continue
            nid = f"z{len(nodes)}"
            nodes[key] = nid
            shape = ", shape=folder" if depth < len(parts) else ""   # folders vs files
            lines.append(f"  {nid} [label={q(parts[depth - 1])}{shape}];")
            lines.append(f"  {nodes['/'.join(parts[:depth - 1])]} -> {nid};")   # arrow from parent
    lines.append("}")
    return "\n".join(lines)


def analyze_zip(data):
    outs = []
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        infos = [i for i in z.infolist() if not i.is_dir()]   # files only, not folders
        outs.append(Output("info", {"Files inside": len(infos),
                                    "Total uncompressed": human_size(sum(i.file_size for i in infos))},
                           "Archive details"))
        # A list of dicts -> each dict becomes a table row
        outs.append(Output("table", pd.DataFrame([{"File": i.filename, "Size": human_size(i.file_size),
                                                   "Compressed": human_size(i.compress_size)} for i in infos[:500]]),
                           "Contents"))
        types = collections.Counter(get_ext(i.filename) or "(none)" for i in infos).most_common(12)
        if types:
            labels, values = zip(*types)
            outs.append(Output("chart", bar_chart(list(labels), list(values), "Files by type", "files"), "File types"))
        outs.append(Output("graphviz", folder_tree_dot([i.filename for i in infos]), "Folder structure"))

        # Preview a few pictures and up to 3 small text files inside the zip
        shown_img = shown_txt = 0               # set both to 0
        for i in infos:
            ext = get_ext(i.filename)
            if ext in IMAGE_EXT and shown_img < MAX_EMBEDDED_IMAGES:
                img = raw_image_to_display(z.read(i))
                if img:
                    outs.append(Output("image", img, i.filename)); shown_img += 1
            elif shown_txt < 3 and i.file_size < 200_000:
                raw = z.read(i)
                if not looks_binary(raw):
                    outs.append(Output("code", truncate(decode_text(raw)[0], 20_000), f"Preview: {i.filename}",
                                       {"language": CODE_LANG.get(ext)}))
                    shown_txt += 1
    return outs


def hex_dump(data, n=512):
    """Show raw bytes like programmers do:  offset | bytes in hexadecimal | readable characters"""
    rows = []
    for off in range(0, min(len(data), n), 16):          # range(start, stop, step) -> 0, 16, 32 ...
        chunk = data[off:off + 16]
        hex_part = " ".join(f"{b:02x}" for b in chunk)    # {b:02x} = number as 2-digit hexadecimal
        ascii_part = "".join(chr(b) if 32 <= b < 127 else "." for b in chunk)   # chr(65) -> "A"
        rows.append(f"{off:08x}  {hex_part:<47}  {ascii_part}")              # :<47 = pad to 47 wide
    return "\n".join(rows)


def analyze_binary(data):
    """Unknown binary file: still 'considered as text' by pulling out readable words."""
    sample = np.frombuffer(data[:5_000_000], dtype=np.uint8)   # bytes -> numbers 0-255
    counts = np.bincount(sample, minlength=256)               # how often each byte value appears
    p = counts[counts > 0] / max(len(sample), 1)
    # Entropy: how "random" the data is. ~8 = compressed or encrypted, lower = more structure.
    entropy = float(-(p * np.log2(p)).sum())
    outs = [Output("info", {"Entropy (bits/byte)": f"{entropy:.2f}  (8 = random/compressed/encrypted)"},
                   "Binary details")]

    # Readable strings: runs of 4+ printable characters (same idea as the Linux "strings" command)
    strings = [s.decode("ascii") for s in re.findall(rb"[\x20-\x7e]{4,}", data[:5_000_000])][:400]
    if strings:
        outs += text_outputs("\n".join(strings), "Readable text found inside the file")
    outs.append(Output("code", hex_dump(data), "Hex dump (first 512 bytes)", {"language": "text"}))

    fig, ax = plt.subplots(figsize=(8, 3))
    ax.bar(range(256), counts, color="#4C78A8", width=1.0)
    ax.set_title("Byte value distribution"); ax.set_xlabel("byte value (0–255)"); ax.set_ylabel("count")
    fig.tight_layout()
    outs.append(Output("chart", fig, "Byte histogram"))
    return outs


# =====================================================================
# MAIN ENTRY POINT - the only function the front-ends need to call
# =====================================================================
def analyze(name, data):
    """Decide what kind of file this is, send it to the right analyze_xxx()
    function, and return ALL the outputs as one list."""
    ext = get_ext(name)                         # from the file NAME, e.g. ".pdf"
    kind = sniff(data)                          # from the file CONTENT, e.g. "pdf"
    guessed = mimetypes.guess_type(name)[0]
    outs = [Output("info", {"File name": name, "Extension": ext or "(none)",
                            "Detected type": guessed or kind or ("text" if not looks_binary(data) else "binary"),
                            "Size": human_size(len(data))}, "File details")]
    if not data:
        outs.append(Output("warning", "The file is empty."))
        return outs

    try:
        # One big if/elif chain: the FIRST matching condition wins.
        # Order matters: e.g. .svg must be checked before general images.
        if ext == ".svg":
            outs += analyze_svg(data)
        elif ext in IMAGE_EXT or kind == "image":
            outs += analyze_image(data)
        elif ext == ".pdf" or kind == "pdf":
            outs += analyze_pdf(data)
        elif ext == ".docx":
            outs += analyze_docx(data)
        elif ext == ".pptx":
            outs += analyze_pptx(data)
        elif ext in (".xlsx", ".xlsm", ".xls"):
            outs += analyze_excel(data)
        elif ext in (".csv", ".tsv"):
            outs += analyze_csv(data, ext)
        elif ext == ".ipynb":
            outs += analyze_notebook(data)
        elif ext == ".json":
            outs += analyze_json(data)
        elif ext in (".md", ".markdown"):
            outs += analyze_markdown(data)
        elif ext in (".html", ".htm"):
            outs += analyze_html(data)
        elif ext == ".xml":
            outs += analyze_xml(data)
        elif ext in AUDIO_EXT or kind == "audio":
            outs += analyze_media(data, ext, "audio")
        elif ext in VIDEO_EXT or kind == "video":
            outs += analyze_media(data, ext, "video")
        elif ext in (".zip", ".jar", ".apk", ".epub") or kind == "zip":
            outs += analyze_zip(data)
        elif ext in GRAPHVIZ_EXT:
            text, _ = decode_text(data)
            outs += [Output("graphviz", text, "Graphviz diagram")] + text_outputs(text, "DOT source", "text", False)
        elif ext in MERMAID_EXT:
            text, _ = decode_text(data)
            outs += [Output("mermaid", text, "Mermaid diagram")] + text_outputs(text, "Mermaid source", "text", False)
        elif not looks_binary(data):
            outs += analyze_text(data, ext, name)          # any other text: .txt, .log, code ...
        else:
            outs += analyze_binary(data)                   # anything else: .exe, .bin, .dat ...
    except Exception as e:
        # Safety net: if ANYTHING above crashes (corrupt file, wrong extension...),
        # we don't crash the app - we explain and fall back to showing it as text.
        outs.append(Output("warning", f"Couldn't fully read this as '{ext or 'unknown'}' ({type(e).__name__}: {e}). "
                                      "Showing it as text instead."))
        outs += analyze_text(data, ext, name) if not looks_binary(data) else analyze_binary(data)
    return outs

# ╔══ NEXT: CELL 3 → create a NEW cell, its first line must be: %%writefile app.py ══╗
# (This comment line is saved inside file_analyzer.py too - that's harmless.)



%%writefile app.py
# ↑ saves this cell into app.py - the Streamlit web page.
# Remember: Streamlit re-runs this WHOLE file from top to bottom every time the user
# does something (e.g. uploads a file).

import matplotlib.pyplot as plt
import pandas as pd
import streamlit as st                          # "st" is the standard nickname
import streamlit.components.v1 as components    # lets us show custom HTML (Mermaid, web pages)

# Import the functions we wrote in file_analyzer.py. Python finds the file because
# it's in the same folder as app.py.
from file_analyzer import (analyze, group_outputs, capabilities, mermaid_html,
                           svg_html, human_size, SECTION_ORDER)

# An emoji for each tab
ICONS = {"Overview": "ℹ️", "Text": "📝", "Visuals": "🖼️", "Charts": "📊",
         "Diagrams": "🧭", "Data": "🗃️", "Preview": "▶️"}


def render(o):
    """Show ONE Output object using the matching Streamlit command.
    Think of it as a translator: Output kind -> Streamlit widget."""
    if o.title:
        st.markdown(f"**{o.title}**")           # **text** = bold in markdown
    k = o.kind
    if k == "info":
        # turn the dict into a 2-column table: Property | Value
        st.dataframe(pd.DataFrame({"Property": list(o.data.keys()),
                                   "Value": [str(v) for v in o.data.values()]}), hide_index=True)
    elif k == "note":
        st.info(o.data)                         # blue box
    elif k == "warning":
        st.warning(o.data)                      # yellow box
    elif k == "text":
        with st.container(height=400):          # a scrollable box 400 pixels tall
            st.text(o.data)
    elif k == "code":
        with st.container(height=400):
            st.code(o.data, language=o.extra.get("language") or None)   # coloured code
    elif k == "markdown":
        with st.container(height=500):
            st.markdown(o.data)
    elif k == "image":
        st.image(o.data)
    elif k == "svg":
        # unsafe_allow_html=True lets Streamlit show our raw HTML <img> tag
        st.markdown(svg_html(o.data), unsafe_allow_html=True)
    elif k == "chart":
        st.pyplot(o.data)                       # show a matplotlib figure
        plt.close(o.data)                       # free the memory it used
    elif k == "table":
        try:
            st.dataframe(o.data)
        except Exception:                       # some mixed-type columns can't be shown...
            st.dataframe(o.data.astype(str))    # ...so convert everything to text and retry
    elif k == "json":
        st.json(o.data)                         # collapsible JSON viewer
    elif k == "graphviz":
        try:
            st.graphviz_chart(o.data)           # draws DOT diagrams in the browser
        except Exception as e:
            st.warning(f"Diagram couldn't be drawn: {e}")
            st.code(o.data)
    elif k == "mermaid":
        components.html(mermaid_html(o.data), height=480, scrolling=True)
    elif k == "html":
        components.html(o.data, height=500, scrolling=True)   # shown in a safe, separate frame
    elif k == "audio":
        st.audio(o.data, format=o.extra.get("mime", "audio/mpeg"))
    elif k == "video":
        st.video(o.data, format=o.extra.get("mime", "video/mp4"))
    st.write("")                                # an empty line for spacing


def main():
    # Page settings - must be the first Streamlit command
    st.set_page_config(page_title="Universal File Viewer", page_icon="🗂️", layout="wide")
    st.title("🗂️ Universal File Viewer")
    st.caption("Upload any file. It's read as text, and shown as images, charts, diagrams, "
               "tables or media wherever possible.")

    # Everything inside "with st.sidebar:" appears in the left sidebar
    with st.sidebar:
        st.header("What it can show")
        st.markdown(
            "- **Text & code**: any text file, with stats and word charts\n"
            "- **Images**: preview, EXIF, colour histogram, OCR text\n"
            "- **PDF**: page images, text, embedded pictures\n"
            "- **Word / PowerPoint**: text, outline diagram, tables, images\n"
            "- **CSV / Excel / JSON**: tables, statistics, charts\n"
            "- **Diagrams**: Graphviz (.dot) and Mermaid (.mmd) files, or diagrams inside markdown\n"
            "- **Python**: code-structure diagram\n"
            "- **Audio / video**: players\n"
            "- **ZIP**: contents, folder tree, previews\n"
            "- **Anything else**: readable strings, hex dump, byte chart"
        )
        st.subheader("Optional features")
        for name, ok in capabilities().items():
            st.write(("✅ " if ok else "❌ ") + name)

    # No "type=" argument -> the uploader accepts ANY file type.
    files = st.file_uploader("Upload one or more files (any type)", accept_multiple_files=True)
    if not files:                               # empty list -> nothing uploaded yet
        st.info("⬆️ Upload a file to begin.")
        return                                  # stop main() here

    for f in files:                             # handle each uploaded file
        st.divider()                            # horizontal line
        data = f.getvalue()                     # the file's content as bytes
        st.header(f"📄 {f.name}  ·  {human_size(len(data))}")
        with st.spinner("Analysing file..."):   # loading animation while analyze() runs
            outputs = analyze(f.name, data)     # ← the "brain" does all the work

        groups = group_outputs(outputs)
        sections = [s for s in SECTION_ORDER if groups[s]]   # only sections that have something
        tabs = st.tabs([f"{ICONS[s]} {s} ({len(groups[s])})" for s in sections])
        for tab, s in zip(tabs, sections):      # pair each tab with its section name
            with tab:                           # everything inside goes into this tab
                for o in groups[s]:
                    render(o)


# This is True only when the file is run directly (streamlit run app.py),
# not when another file imports it. It's a standard Python pattern.
if __name__ == "__main__":
    main()

# ╔══ NEXT: CELL 4a → NEW cell, first line: %%writefile requirements.txt ══╗



%%writefile requirements.txt
streamlit
pandas
numpy
matplotlib
pillow
pymupdf
python-docx
python-pptx
openpyxl
chardet
pytesseract
# ↑ requirements.txt lists the Python packages Streamlit Cloud must install.
#   Lines starting with # are ignored by pip, so these notes are safe here.
# ╔══ NEXT: CELL 4b → NEW cell, first line: %%writefile packages.txt  (no comments in that file!) ══╗



%%writefile packages.txt
tesseract-ocr



# ╔═════════════ CELL 5 : view files DIRECTLY in the Colab output ═════════════╗
# (A normal Python cell - comments at the top are fine here.)
# packages.txt above lists SYSTEM programs for Streamlit Cloud (like apt-get in CELL 1).
# Keep that file to package names only - no comment lines.

import base64, html, importlib, io, json        # several imports on one line is allowed
import matplotlib.pyplot as plt
import pandas as pd
from IPython.display import display, HTML, Markdown   # tools to show things in notebook output
from PIL import Image
from google.colab import files                  # Colab's upload button

import file_analyzer
importlib.reload(file_analyzer)                 # re-load the file so your edits are picked up
from file_analyzer import analyze, group_outputs, mermaid_html, svg_html, SECTION_ORDER, MAX_MEDIA_MB

try:
    import graphviz                             # draws DOT diagrams inside Colab
except ImportError:
    graphviz = None


def scroll_box(text):
    """A scrollable grey box for long text. html.escape stops text like '<b>' being treated as HTML."""
    return HTML(f"<pre style='max-height:400px;overflow:auto;white-space:pre-wrap;"
                f"border:1px solid #ddd;padding:8px'>{html.escape(text)}</pre>")


def show_colab(o):
    """Same job as render() in app.py, but for the Colab notebook output."""
    if o.title:
        display(HTML(f"<b>{html.escape(o.title)}</b>"))
    k = o.kind
    if k == "info":
        display(pd.DataFrame({"Property": list(o.data.keys()), "Value": [str(v) for v in o.data.values()]}))
    elif k in ("note", "warning"):
        colour = "#E8F4FD" if k == "note" else "#FFF4E5"     # blue for notes, orange for warnings
        display(HTML(f"<div style='background:{colour};padding:8px;border-radius:6px'>{html.escape(o.data)}</div>"))
    elif k in ("text", "code"):
        display(scroll_box(o.data))
    elif k == "markdown":
        display(Markdown(o.data))
    elif k == "image":
        fmt = (Image.open(io.BytesIO(o.data)).format or "png").lower()   # "png", "jpeg", "gif"...
        b64 = base64.b64encode(o.data).decode()
        display(HTML(f'<img src="data:image/{fmt};base64,{b64}" style="max-width:100%;max-height:600px">'))
    elif k == "svg":
        display(HTML(svg_html(o.data)))
    elif k == "chart":
        display(o.data)                         # Colab knows how to draw matplotlib figures
        plt.close(o.data)
    elif k == "table":
        display(o.data)                         # pandas tables display nicely in Colab
    elif k == "json":
        # json.dumps(..., indent=2) = pretty-printed JSON text
        display(scroll_box(json.dumps(o.data, indent=2, ensure_ascii=False, default=str)[:50_000]))
    elif k == "graphviz":
        try:
            display(graphviz.Source(o.data))
        except Exception as e:
            display(HTML(f"<i>Diagram couldn't be drawn ({html.escape(str(e))}); DOT source:</i>"))
            display(scroll_box(o.data))
    elif k == "mermaid":
        display(HTML(mermaid_html(o.data)))
    elif k == "html":
        # srcdoc + sandbox = show the uploaded page in a safe, separate frame
        display(HTML(f'<iframe srcdoc="{html.escape(o.data)}" sandbox '
                     f'style="width:100%;height:450px;border:1px solid #ccc"></iframe>'))
    elif k in ("audio", "video"):
        if len(o.data) > MAX_MEDIA_MB * 1024 * 1024:          # MB -> bytes
            display(HTML(f"<i>{k.title()} is larger than {MAX_MEDIA_MB} MB; use the Streamlit app to play it.</i>"))
        else:
            b64 = base64.b64encode(o.data).decode()
            tag = "audio" if k == "audio" else "video style='max-width:100%'"
            display(HTML(f"<{tag} controls src='data:{o.extra['mime']};base64,{b64}'></{k}>"))


# files.upload() shows a "Choose Files" button and returns a dict: {file name: file bytes}
uploaded = files.upload()
for name, data in uploaded.items():
    display(HTML(f"<hr><h2>📄 {html.escape(name)}</h2>"))
    groups = group_outputs(analyze(name, data))         # same "brain" as the Streamlit app
    for section in SECTION_ORDER:
        if groups[section]:
            display(HTML(f"<h3 style='color:#2E86C1'>{section}</h3>"))
            for o in groups[section]:
                show_colab(o)



# ╔═════════════ CELL 6 : run the Streamlit app from Colab (public link) ═════════════╗
# Colab runs on Google's computer, so "localhost" isn't reachable from your browser.
# We start Streamlit there, then use a free Cloudflare "tunnel" to get a public link.
import os, re, subprocess, time

# 1) Start Streamlit in the BACKGROUND. subprocess.Popen runs a program without waiting
#    for it to finish. Its messages are saved in streamlit.log (open it if something fails).
subprocess.Popen(["streamlit", "run", "app.py", "--server.port", "8501", "--server.headless", "true"],
                 stdout=open("streamlit.log", "w"), stderr=subprocess.STDOUT)
time.sleep(5)                                   # wait 5 seconds for it to start

# 2) Download the cloudflared program once (skipped if it's already there)
if not os.path.exists("cloudflared"):
    !wget -q https://github.com/cloudflare/cloudflared/releases/latest/download/cloudflared-linux-amd64 -O cloudflared
    !chmod +x cloudflared

# 3) Open the tunnel and read its messages line by line until we find the public link
tunnel = subprocess.Popen(["./cloudflared", "tunnel", "--url", "http://localhost:8501"],
                          stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True)
for line in tunnel.stdout:
    match = re.search(r"https://[-a-z0-9]+\.trycloudflare\.com", line)
    if match:
        print("✅ Your app is live at:", match.group(0))
        print("   (keep this cell running; stop it to shut the app down)")
        break                                   # stop reading - we found the link
